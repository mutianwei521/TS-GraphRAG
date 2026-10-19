"""Scientific parity checks for the public Source Data and derived Table S17."""
import csv
import importlib.util
import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("source_builder", ROOT / "scripts/build_source_data.py")
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


def data_sheet(name):
    payload = json.loads((ROOT / f"source_data/sheets/{name}.json").read_text(encoding="utf-8"))
    return pd.DataFrame(payload["rows"][1:], columns=payload["rows"][0])


def test_all_36_sheets_rebuild_with_native_cells_and_s17(tmp_path):
    manifest, sheets = builder.load_sheets()
    assert manifest["canonical_sha256"] == "cbff510f2bf19abce1b65c8b75694ee14db8fb30606a0d8205fd02d3ef913a52"
    assert len(sheets) == 36
    report = builder.build(tmp_path)
    assert report["table_s17_recomputed_from_test_json"]
    rebuilt = openpyxl.load_workbook(tmp_path / "Source_Data.xlsx", data_only=False)
    assert rebuilt.sheetnames == [sheet["name"] for sheet in sheets]
    for sheet in sheets:
        result = rebuilt[sheet["name"]]
        assert list(result.values) == [tuple(row) for row in sheet["rows"]]
    assert rebuilt["TableS17_node_holdout"].max_row == 14
    assert all(c.value is None for c in rebuilt["TableS17_node_holdout"][11])
    assert rebuilt["TableS17_node_holdout"]["A12"].value == "Notes"
    assert rebuilt["TableS17_node_holdout"]["A14"].value == "Comparison"
    # Text node IDs survive numeric-looking CSV cells.
    assert isinstance(rebuilt["Sensor_allocations"]["E2"].value, str)


def test_e4_repetition_values_match_the_numerical_workbook():
    saved = data_sheet("Fig4abc_per_run")
    original = pd.read_csv(ROOT / "source_data/figure_inputs/results/E4_mean_std_20260909/all_run_accuracies.csv")
    keys = ["panel", "protocol", "method", "run_seed"]
    a = saved.sort_values(keys).reset_index(drop=True)
    b = original.sort_values(keys).reset_index(drop=True)
    assert len(a) == len(b) == 675
    assert a[keys].equals(b[keys])
    np.testing.assert_allclose(a.accuracy_percent, b.accuracy_percent, rtol=0, atol=1e-12)
    # The plotted destination observations account for the other 15 run cells.
    assert len(data_sheet("Fig4d_destination_per_run")) == 15


def test_sensitivity_34_settings_and_paired_seed_summary():
    raw = data_sheet("Sensitivity_per_seed")
    summary = data_sheet("FigS6_TableS16")
    assert len(raw) == 680 and len(summary) == 136
    assert raw.setting.nunique() == 34
    assert set(raw.seed) == {41, 141, 241, 341, 441}
    assert set(raw.protocol) == {"H0", "H2", "H4", "H8"}
    grouped = raw.groupby(["protocol", "setting"]).top1_pct.agg(["mean", "std"])
    published = summary.set_index(["protocol", "setting"]).sort_index()
    np.testing.assert_allclose(grouped["mean"], published.top1_pct_mean, rtol=0, atol=1e-12)
    np.testing.assert_allclose(grouped["std"], published.top1_pct_sd, rtol=0, atol=1e-12)
    defaults = raw[raw.setting.eq("default")].set_index(["seed", "protocol"]).top1_pct
    for (protocol, setting), frame in raw.groupby(["protocol", "setting"]):
        changes = frame.set_index("seed").top1_pct - defaults.xs(protocol, level="protocol")
        expected = published.loc[(protocol, setting), "top1_pct_mean"] - defaults.xs(protocol, level="protocol").mean()
        assert abs(changes.mean() - expected) < 1e-12


def test_s17_derives_test_percentages_and_paired_p_without_rounding_early():
    rows, numeric = builder.table_s17(ROOT / "data/frozen/node_selection")
    assert len(rows) == 9
    assert rows[0] == ["H0", 85.7, 86.67, 85.7, 85.78, 1.0]
    assert rows[8] == ["H8", 65.33, 46.3, 65.33, 65.48, 0.5]
    assert numeric[0]["projection_top1_pct"] == 100 * (1157 / 1350)
    assert numeric[8]["masked_projection_top1_pct"] == 100 * (884 / 1350)
    table = data_sheet("TableS17_node_holdout")
    assert table.iloc[:9].values.tolist() == rows


def test_current_si_snapshots_have_weights_and_correct_s16_s17_content():
    table_dir = ROOT / "source_data/manuscript_tables"
    manifest = json.loads((ROOT / "source_data/manuscript_table_manifest.json").read_text(encoding="utf-8"))
    assert manifest["table_count"] == 28
    weights = pd.read_csv(table_dir / "TableS1.csv")
    np.testing.assert_allclose(weights.iloc[:, 2:].sum().values, [1, 1, 1], atol=1e-15)
    sensitivity = pd.read_csv(table_dir / "TableS16.csv")
    holdout = pd.read_csv(table_dir / "TableS17.csv")
    assert sensitivity.columns[0] == "Setting" and len(sensitivity) == 34
    assert holdout.columns[0] == "Protocol" and len(holdout) == 9


def test_figure_map_routes_every_current_panel_and_keeps_attribution_separate():
    with (ROOT / "reproducibility/figure_table_map.csv").open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    for figure in [*[f"Fig{i}" for i in range(1, 8)], *[f"FigS{i}" for i in range(1, 7)]]:
        assert any(row["item"] == figure and row["panel"] == "all" for row in rows)
        assert any(row["item"] == figure and row["kind"] == "data_panel" for row in rows)
    for row in rows:
        assert all(row[key] for key in ["command", "input", "output", "expected", "status"])
        assert "C:/" not in row["command"] and "C:\\" not in row["command"]
    s13 = [row for row in rows if row["item"].startswith("TableS13")]
    assert len(s13) == 2 and all("Hydraulic_summary" in row["input"] for row in s13)
    attribution = next(row for row in rows if row["item"] == "SupplementaryNote1.9")
    assert "Strict_H9_attribution" in attribution["input"]
    standalone = next(row for row in rows if row["item"] == "FigS6_standalone_20261017")
    assert standalone["version"] == "standalone_20261017_not_integrated"
