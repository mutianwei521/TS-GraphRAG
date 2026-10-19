"""Build the complete, public Source Data workbook with openpyxl.

Normal use needs only the files in this repository. --import-workbook is a
maintainer operation that exports a reviewed workbook to rectangular CSV and
typed JSON. JSON is authoritative for cell types; CSV is a convenient view.
No worksheet rows are discarded and only row 1 is designated as a header.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source_data"
EXPECTED_SHEETS = 36


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, allow_nan=False,
                               indent=2) + "\n", encoding="utf-8")


def csv_text(rows) -> str:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows(rows)
    return buffer.getvalue()


def export_workbook(path: Path, destination: Path) -> dict:
    """Export all native values, cell types, dimensions, and merged ranges."""
    workbook = openpyxl.load_workbook(path, data_only=False)
    if len(workbook.sheetnames) != EXPECTED_SHEETS:
        raise ValueError(f"Expected {EXPECTED_SHEETS} sheets; found {len(workbook.sheetnames)}")
    records = []
    for sheet in workbook:
        rows = []
        cell_types = []
        formats = {}
        for row in sheet.iter_rows():
            values, types = [], []
            for cell in row:
                value = cell.value
                if isinstance(value, (dt.date, dt.datetime, dt.time)):
                    value = {"iso8601": value.isoformat(), "kind": type(value).__name__}
                values.append(value)
                types.append(cell.data_type)
                if cell.number_format != "General":
                    formats[cell.coordinate] = cell.number_format
            rows.append(values)
            cell_types.append(types)
        name = sheet.title
        json_path = destination / "sheets" / f"{name}.json"
        csv_path = destination / "sheets" / f"{name}.csv"
        payload = {"name": name, "max_row": sheet.max_row,
                   "max_column": sheet.max_column, "rows": rows,
                   "cell_types": cell_types, "number_formats": formats,
                   "merged_ranges": [str(r) for r in sheet.merged_cells.ranges],
                   "sheet_state": sheet.sheet_state,
                   "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                   "header_rows": [1],
                   "row_policy": "All rows within worksheet dimensions, including blank and note rows, are retained.",
                   "column_widths": {k: v.width for k, v in sheet.column_dimensions.items()},
                   "row_heights": {str(k): v.height for k, v in sheet.row_dimensions.items() if v.height}}
        write_json(json_path, payload)
        csv_path.write_text(csv_text(rows), encoding="utf-8", newline="")
        records.append({"name": name, "rows_including_header": sheet.max_row,
                        "columns": sheet.max_column,
                        "json": json_path.relative_to(destination).as_posix(),
                        "csv": csv_path.relative_to(destination).as_posix(),
                        "json_sha256": sha256(json_path), "csv_sha256": sha256(csv_path)})
    manifest = {"schema_version": 1, "sheet_count": len(records),
                "canonical_workbook": "0829/reviewer_closure_20261010/Source_Data.xlsx",
                "canonical_sha256": sha256(path),
                "canonical_bytes": path.stat().st_size,
                "value_contract": "Typed JSON preserves native cell values and data types. CSV includes row 1 and every blank/note row; empty CSV fields are disambiguated by JSON. Workbook presentation styles are not part of numerical parity.",
                "sheets": records}
    write_json(destination / "workbook_manifest.json", manifest)
    return manifest


def load_sheets(source: Path = SOURCE) -> tuple[dict, list[dict]]:
    manifest = json.loads((source / "workbook_manifest.json").read_text(encoding="utf-8"))
    if manifest["sheet_count"] != EXPECTED_SHEETS or len(manifest["sheets"]) != EXPECTED_SHEETS:
        raise ValueError("Incomplete Source Data manifest")
    sheets = []
    for entry in manifest["sheets"]:
        path = source / entry["json"]
        if sha256(path) != entry["json_sha256"] or sha256(source / entry["csv"]) != entry["csv_sha256"]:
            raise ValueError(f"Changed frozen source: {entry['name']}")
        sheet = json.loads(path.read_text(encoding="utf-8"))
        if len(sheet["rows"]) != entry["rows_including_header"]:
            raise ValueError(f"Row-count mismatch: {entry['name']}")
        if any(len(row) != entry["columns"] for row in sheet["rows"]):
            raise ValueError(f"Nonrectangular sheet: {entry['name']}")
        expected_csv = csv_text(sheet["rows"])
        if (source / entry["csv"]).read_text(encoding="utf-8") != expected_csv:
            raise ValueError(f"CSV/JSON disagreement: {entry['name']}")
        sheets.append(sheet)
    return manifest, sheets


def table_s17(selection: Path) -> tuple[list[list], list[dict]]:
    """Recompute the rounded publication cells and retain unrounded percentages."""
    ordinary = json.loads((selection / "test_results.json").read_text(encoding="utf-8"))
    masked = json.loads((selection / "test_masked_results.json").read_text(encoding="utf-8"))
    base = {r["protocol"]: r for r in ordinary}
    mask = {r["protocol"]: r for r in masked}
    protocols = [f"H{i}" for i in range(9)]
    if sorted(base) != protocols or sorted(mask) != protocols:
        raise ValueError("Table S17 requires exactly H0-H8")
    rows, numeric = [], []
    for protocol in protocols:
        b, m = base[protocol], mask[protocol]
        values = [100 * b["test_top1_projection_recomputed"],
                  100 * b["test_top1_default_fusion"],
                  100 * b["test_top1_saved_cosine"],
                  100 * m["test_top1_frozen"]]
        if b["test_top1_projection_recomputed"] != m["test_top1_projection_tau0"]:
            raise ValueError(f"Inconsistent unmasked projection: {protocol}")
        rows.append([protocol, *[round(v, 2) for v in values],
                     m["mcnemar_frozen_vs_projection_p"]])
        numeric.append(dict(zip(["protocol", "projection_top1_pct", "default_fusion_top1_pct",
                                 "cosine_top1_pct", "masked_projection_top1_pct", "mcnemar_p"],
                                [protocol, *values, m["mcnemar_frozen_vs_projection_p"]])))
    return rows, numeric


def build(output: Path, source: Path = SOURCE,
          selection: Path = ROOT / "data/frozen/node_selection") -> dict:
    manifest, sheets = load_sheets(source)
    derived_rows, numeric = table_s17(selection)
    table = next(s for s in sheets if s["name"] == "TableS17_node_holdout")
    if table["rows"][1:10] != derived_rows:
        raise ValueError("Table S17 frozen cells do not match the unrounded test results")
    table["rows"][1:10] = derived_rows
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for payload in sheets:
        sheet = workbook.create_sheet(payload["name"])
        for i, values in enumerate(payload["rows"], 1):
            for j, value in enumerate(values, 1):
                if isinstance(value, dict) and "iso8601" in value:
                    value = getattr(dt, value["kind"]).fromisoformat(value["iso8601"])
                cell = sheet.cell(i, j, value)
                cell.data_type = payload["cell_types"][i-1][j-1]
        for coordinate, number_format in payload["number_formats"].items():
            sheet[coordinate].number_format = number_format
        for merged in payload["merged_ranges"]:
            sheet.merge_cells(merged)
        for column, width in payload["column_widths"].items():
            sheet.column_dimensions[column].width = width
        for row, height in payload["row_heights"].items():
            sheet.row_dimensions[int(row)].height = height
        sheet.sheet_state = payload["sheet_state"]
        sheet.freeze_panes = payload["freeze_panes"]
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315A76")
    output.mkdir(parents=True, exist_ok=True)
    workbook_path = output / "Source_Data.xlsx"
    workbook.save(workbook_path)
    # Reopen the file actually written, checking native values and blank rows.
    rebuilt = openpyxl.load_workbook(workbook_path, data_only=False)
    if rebuilt.sheetnames != [s["name"] for s in sheets]:
        raise ValueError("Rebuilt worksheet order changed")
    for payload in sheets:
        sheet = rebuilt[payload["name"]]
        actual = [[sheet.cell(i, j).value for j in range(1, payload["max_column"] + 1)]
                  for i in range(1, payload["max_row"] + 1)]
        if actual != payload["rows"]:
            raise ValueError(f"Rebuilt native-cell mismatch: {sheet.title}")
    numeric_path = output / "TableS17_numeric.csv"
    with numeric_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(numeric[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(numeric)
    report = {"status": "passed", "sheet_count": EXPECTED_SHEETS,
              "native_cell_parity": True, "blank_and_note_rows_retained": True,
              "table_s17_recomputed_from_test_json": True,
              "canonical_workbook_sha256": manifest["canonical_sha256"],
              "workbook_sha256": sha256(workbook_path),
              "selection_input_sha256": {name: sha256(selection / name) for name in
                                          ["test_results.json", "test_masked_results.json"]}}
    write_json(output / "build_report.json", report)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "results/source_data")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--selection", type=Path, default=ROOT / "data/frozen/node_selection")
    parser.add_argument("--import-workbook", type=Path,
                        help="Maintainer export; requires the canonical 36-sheet workbook")
    args = parser.parse_args()
    if args.import_workbook:
        report = export_workbook(args.import_workbook, args.source)
        print(json.dumps({"status": "exported", "sheet_count": report["sheet_count"]}))
    else:
        print(json.dumps(build(args.output, args.source, args.selection), indent=2))


if __name__ == "__main__":
    main()
